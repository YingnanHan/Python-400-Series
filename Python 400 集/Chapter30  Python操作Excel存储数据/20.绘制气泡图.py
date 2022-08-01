import openpyxl
from openpyxl.chart import BubbleChart
from openpyxl.chart import Reference
from openpyxl.chart import Series

#创建工作簿
wb = openpyxl.Workbook()
ws = wb.create_sheet("bubbleChart")
ws.title = "Bubble Chart"

#绘制气泡图
rows = [
    ("Number of products","Sales in USD","Market shares"),
    (14,12200,15),
    (20,60000,33),
    (18,24400,10),
    (22,32000,42),
    (),
    (12,8200,18),
    (15,50000,30),
    (19,22400,15),
    (25,25000,50)
]

#添加数据
for row in rows:
    ws.append(row)

bubble_chart = BubbleChart() #定义bubbleChart对象
bubble_chart.style = 18 #定义气泡图样式
#添加第一组数据
xvalue = Reference(ws,min_col=1,max_col=1,min_row=2,max_row=5)
yvalue = Reference(ws,min_col=2,max_col=2,min_row=2,max_row=5)
size = Reference(ws,min_col=2,max_col=2,min_row=2,max_row=5)
series = Series(values=yvalue,xvalues=xvalue,zvalues=size,title="2013")#创建Series
bubble_chart.series.append(series)

#添加第二组数据
xvalue = Reference(ws,min_col=1,max_col=1,min_row=7,max_row=10)
yvalue = Reference(ws,min_col=2,max_col=2,min_row=7,max_row=10)
size = Reference(ws,min_col=2,max_col=2,min_row=7,max_row=10)
series = Series(values=yvalue,xvalues=xvalue,zvalues=size,title="2014")#创建Series
bubble_chart.series.append(series)

#在工作表中添加气泡图
ws.add_chart(bubble_chart,"F6")

#保存工作簿
wb.save("bubblechart.xlsx")