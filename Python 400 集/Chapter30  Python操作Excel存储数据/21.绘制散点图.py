import openpyxl
from openpyxl.chart import (BarChart, Reference, ScatterChart,Series)
#绘制散点图
rows = [
            ['Size', 'Batch 1', 'Batch 2'],
            [2, 40, 30],
            [3, 40, 25],
            [4, 50, 30],
            [5, 30, 25],
            [6, 25, 35],
            [7, 20, 40],
]

#创建工作簿
wb = openpyxl.Workbook()
ws = wb.create_sheet('Scatter Chart')
ws.title = "Scatter Chart"

#添加数据
for row in rows:
    ws.append(row)

chart = ScatterChart()#定义散点图对象

chart.title = '散点图' #设置散点图样式
chart.style = 13
chart.x_axis.title = 'Size'
chart.y_axis.title = 'Percentage'

#使用一个循环将两条折线同时画出
xvalues = Reference(ws, min_col=1, min_row=2, max_row=7)
for c in range(2,4):
    yvalues = Reference(ws,min_col=c,min_row=1,max_row=7)
    series = Series(values=yvalues,xvalues=xvalues,title_from_data=True)
    chart.series.append(series)

ws.add_chart(chart,'E1')
wb.save('scatterchart.xlsx')